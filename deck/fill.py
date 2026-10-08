import copy, sys
from pptx import Presentation
from pptx.oxml.ns import qn

prs = Presentation(sys.argv[1])
S = list(prs.slides)

def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    raise KeyError(idx)

def fill(shape, paras):
    """paras: list of str | (str, level) | list-of-runs [(text, run_index_to_copy)]"""
    txBody = shape.text_frame._txBody
    ps = txBody.findall(qn('a:p'))
    tmpl = copy.deepcopy(ps[0]) if ps else txBody.makeelement(qn('a:p'), {})
    # collect run templates from all original runs
    runs_t = [copy.deepcopy(r) for p in ps for r in p.findall(qn('a:r'))]
    for p in ps:
        txBody.remove(p)
    for item in paras:
        level = 0
        if isinstance(item, tuple) and isinstance(item[1], int) and not isinstance(item[0], tuple):
            text, level = item; runs = [(text, 0)]
        elif isinstance(item, list):
            runs = item
        else:
            runs = [(item, 0)]
        p = copy.deepcopy(tmpl)
        for child in list(p):
            if child.tag != qn('a:pPr'):
                p.remove(child)
        pPr = p.find(qn('a:pPr'))
        if level:
            if pPr is None:
                pPr = p.makeelement(qn('a:pPr'), {}); p.insert(0, pPr)
            pPr.set('lvl', str(level))
        elif pPr is not None and 'lvl' in pPr.attrib:
            del pPr.attrib['lvl']
        for text, ri in runs:
            if runs_t:
                r = copy.deepcopy(runs_t[min(ri, len(runs_t)-1)])
                for c in list(r):
                    if c.tag not in (qn('a:rPr'),):
                        r.remove(c)
                t = r.makeelement(qn('a:t'), {}); t.text = text; r.append(t)
            else:
                r = p.makeelement(qn('a:r'), {}); t = r.makeelement(qn('a:t'), {}); t.text = text; r.append(t)
            p.append(r)
        txBody.append(p)

def ensure_ph(slide):
    if len(slide.placeholders) == 0:
        slide.shapes.clone_layout_placeholders(slide.slide_layout)

C = {}
exec(open(sys.argv[3]).read())
prs.save(sys.argv[2])
